"""Таблица эталона для согласования с Алексеем: build/etalon_soglasovanie.docx. Без персональных данных (обезличено кодом).
01.10.2026: обращение на «Вы»; K041 и P5–P8 исключены; K059 → тип 1; общий вопрос по группе K020/K021/K024/K046;
вопрос по K030 с учётом регламента (OPIS п. 10); просьба о втором разборе.
Правки Вадима 01.10: претензии законченной фразой, СП 31-105 без пункта, как у Алексея, K059 — вопрос «тип 1?», абзац об исключённых."""
import json, re, sys
sys.path.insert(0, "/opt/captain-sam/src")
from anonymize import anonymize
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EXCLUDE = {"K041", "P5", "P6", "P7", "P8"}
ET = [e for e in json.load(open("/opt/captain-sam/build/etalon.json", encoding="utf-8")) if e["id"] not in EXCLUDE]
VERDICT = {"K059": 1}   # решение Вадима 01.10: опечатка признана → подтверждается (в etalon.json пока не внесено)
TYPES = {1: "1 — подтверждается нормативом", 2: "2 — опровергается нормативом / регламентом / практикой", 3: "3 — встречный вопрос",
         4: "4 — вне зоны ответственности", 5: "5 — основания не найдено"}
GROUP = ["K020", "K021", "K024", "K046"]
GROUP_Q = ("Общий вопрос по строкам K020, K021, K024, K046. Клиент требует X, а по стандарту компании X и так делается "
           "(анкеровка обвязки, схема крепления OSB, связь стропил с мауэрлатом, вентзазор). Какой тип ставим: "
           "2 — претензия снимается («делаем по стандарту»), или 1 — норма подтверждает требование клиента? Сейчас стоит 2. "
           "Ответ применим ко всем таким случаям, не только к этим четырём.")
QUESTIONS = {
    "K030": "В регламенте: «крепим по требованиям производителя (способ и количество креплений)». С учётом этого Ваш вердикт "
            "по-прежнему 3 — встречный вопрос (расчёт «слабого крепления»), или 2 — опровергается регламентом?",
    "K059": "Предлагаем тип 1 — претензия обоснована. Согласны?",
    "P9–10": "Вы ответили опровержением (2). Предлагаем эталон 3 — встречный вопрос: при 15 кВт на одной фазе 63–80 А верно; "
             "какая мощность и фазность выделены по ТУ? Подтверждаете 3?",
}
for i in GROUP: QUESTIONS[i] = "См. общий вопрос над таблицей."
ANSWER_BRIEF = {
    "K059": "Опечатка в документе признана — претензия верна.",
    "P9–10": "Три фазы, 15 кВт: ток 22,8 А, автомат 25 А; 32 А позволил бы ~21 кВт — превышение выделенной мощности (СП 256 п. 12.5). "
             "Одна фаза, 5 кВт: ~24 А, автомат 25 А.",
}

# Претензия законченной фразой — без обрывков колонок исходной таблицы («| Актуально |», «Явная», висящие «)» и «,»). Только в docx.
CLAIM_BRIEF = {
    "K023": "Анкеровка каркаса к фундаменту — удержание от опрокидывания.",
    "K057": "Черновой пол OSB 9 мм — тонко, прогибается при шаге лаг 600 мм.",
    "K059": "Брус 38×1900 мм в межэтажном перекрытии — опечатка, нужно уточнить.",
    "K060": "Чистовой пол OSB 22 мм — нужен финишный слой.",
    "K061": "Входная дверь (КНР) без терморазрыва — зимой при −7 °C и ветре будет промерзать.",
    "K062": "Бойлер 50 литров — мало для двухэтажного дома.",
}

def brief(t, n=170):
    t = anonymize(" ".join(t.split()))[0]
    if len(t) <= n: return t
    cut = t[:n]; m = max(cut.rfind(". "), cut.rfind("; "), cut.rfind(" — "))
    return (cut[:m + 1] if m > 60 else cut.rsplit(" ", 1)[0]) + " …"

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_); tcPr.append(s)

d = Document()
sec = d.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, side, Cm(1.5))
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(9)
nk = sum(1 for e in ET if e["id"].startswith("K")); np_ = len(ET) - nk
d.add_heading("Эталон приёмки: пары «претензия → ответ» и тип вердикта — на согласование", level=1)
d.add_paragraph(f"Проект разбора претензий. {len(ET)} пары: {nk} — разбор по объекту в с. Каменское, {np_} — претензия по вводному автомату (пп. 9–10). "
                "Тип вердикта проставлен по Вашим ответам; прошу проверить каждую строку. Строки, выделенные цветом, — спорные, "
                "в последней колонке вопрос. Персональные данные покупателей удалены.")
d.add_paragraph("Типы вердикта: " + " · ".join(TYPES.values()))
p = d.add_paragraph(); r = p.add_run(GROUP_Q); r.bold = True
t = d.add_table(rows=1, cols=6); t.style = "Table Grid"
for i, h in enumerate(["№", "Пара", "Претензия (кратко)", "Ваш ответ (кратко)", "Тип вердикта", "Вопрос для согласования"]):
    c = t.rows[0].cells[i]; c.text = h; c.paragraphs[0].runs[0].bold = True; shade(c, "D9D9D9")
for n, e in enumerate(ET, 1):
    q = QUESTIONS.get(e["id"])
    v = TYPES.get(VERDICT.get(e["id"], e["verdict"]), "— не определён (см. вопрос)")
    row = t.add_row().cells
    vals = [str(n), e["id"], brief(CLAIM_BRIEF.get(e["id"], e["claim"])), anonymize(re.sub(r"СП 31-105 п\. [\d.]+(?:x)?", "СП 31-105-2002", ANSWER_BRIEF.get(e["id"], e.get("note") or "")))[0], v, q or ""]
    for i, val in enumerate(vals): row[i].text = val
    if q:
        for c in row: shade(c, "FCE4D6")
widths = [Cm(0.8), Cm(1.3), Cm(8.5), Cm(7.5), Cm(3.6), Cm(5.5)]
for r in t.rows:
    for i, w in enumerate(widths): r.cells[i].width = w
d.add_paragraph()
d.add_paragraph("Цвет строки — спорная разметка, нужен ответ. Остальные строки — подтвердите или исправьте тип.")
p = d.add_paragraph(); p.add_run("Исключены из приёмки (прошу подтвердить): ").bold = True
p.add_run("K041 сайдинг — коммерческая позиция, в типы вердикта не укладывается, в ответе система даст шаблон „выбор покупателя“; "
          "вопросы 5–8 — это вопросы клиента, а не претензии, по ним система проверяется только на точность ссылок.")
p = d.add_paragraph(); p.add_run("Просьба о втором разборе. ").bold = True
p.add_run("Сейчас эталон построен на одном объекте (Каменское) — этого мало, чтобы честно оценить систему: часть ошибок может быть "
          "особенностью именно этого разбора. Если есть возможность, пришлите, пожалуйста, ещё один разбор претензий "
          "по другому объекту в том же формате «претензия → Ваш ответ».")
out = "/opt/captain-sam/build/etalon_soglasovanie.docx"; d.save(out)
# контроль ПДн: весь текст документа ещё раз через обезличиватель — не должно появиться новых меток
full = "\n".join([pp.text for pp in d.paragraphs] + [c.text for r in t.rows for c in r.cells])
_, found = anonymize(full)
print(out, "| строк:", len(t.rows) - 1, "| спорных:", sum(1 for e in ET if e["id"] in QUESTIONS), "| ПДн-меток при повторной проверке:", len(found),
      "| «Алексе» в тексте:", len(re.findall(r"Алексе", full)))
