"""Проверка нового типа документа ДО чтения текста (правило 7а, 07.10.2026): только счётчики меток обезличивателя
и список слов с заглавной буквы, которых нет в словаре. Текст документа не выводится.
Словарь — слова из индексов норм и регламента (build/normy_index.json, build/reglament_index.json) + ALLOW обезличивателя.
Слово из списка может оказаться фамилией или именем, которое словарь имён не знает, — тогда расширить src/anonymize.py
и прогнать проверку снова. Читать текст — только когда в списке не осталось похожего на ПДн.

  python3 src/pdn_check.py materials/objects/test1-kp.pdf
"""
import os
import collections, json, re, subprocess, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))   # 17.10: свой src — копия ветки не исполняет боевой код
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from anonymize import anonymize, ALLOW

B = f"{ROOT}/build"
WORD = re.compile(r"[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё\-]+")

def read(path):
    if path.endswith(".pdf"):
        return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout
    if path.endswith(".docx"):
        from docx import Document
        d = Document(path)
        return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    if path.endswith(".odt"):   # 19.10: средствами Python — odt2txt на Питере нет, проверка .odt падала
        import zipfile, html
        x = zipfile.ZipFile(path).read("content.xml").decode("utf-8")
        x = re.sub(r"<text:(?:p|h)\b[^>]*>|<text:line-break/>|<table:table-cell\b[^>]*>", "\n", x)
        x = re.sub(r"<text:s(?: text:c=\"(\d+)\")?/>", lambda m: " " * int(m.group(1) or 1), x)
        x = re.sub(r"<text:tab/>", "\t", x)
        return html.unescape(re.sub(r"<[^>]+>", "", x))
    return open(path, encoding="utf-8", errors="replace").read()

def vocab():
    v = {w.lower() for w in ALLOW}
    N = json.load(open(f"{B}/normy_index.json", encoding="utf-8"))
    for d in N.values():
        for p in d["points"].values(): v.update(w.lower() for w in WORD.findall(p["text"]))
    R = json.load(open(f"{B}/reglament_index.json", encoding="utf-8"))
    for x in (R.values() if isinstance(R, dict) else R): v.update(w.lower() for w in WORD.findall(x.get("text", "")))
    return v

text, mapping = anonymize(read(sys.argv[1]))
print("меток:", dict(collections.Counter(k.strip("<>").rsplit("_", 1)[0] for k in mapping)) or 0)
V = vocab()
text = re.sub(r"<[А-ЯA-Z]+_\d+>", " ", text)   # метки обезличивателя — не слова документа
unknown = collections.Counter(w for w in WORD.findall(text) if w[0].isupper() and w.lower() not in V)
print(f"слов с заглавной вне словаря: {len(unknown)}")
for w, n in sorted(unknown.items()): print(f"  {w} ×{n}")
