"""Документы (docx), 09.10.2026 — генератор v2. Форма — предложение, заказчик поправит; шапка без логотипа.
render — «Ответ покупателю»: сквозная нумерация 1, 2, 3; по пункту: замечание → вывод → рассуждение → основание (документ, пункт,
  цитата; регламент = «Стандарт компании») → расчёт → что мы делаем / что нужно уточнить.
render_engineer — «Черновик для инженера» (ВНУТРЕННИЙ): исходный номер пункта, вероятный вердикт, почему не выдано, найденные
  основания с пометками, заготовка ответа по той же структуре и список того, что не подтверждено (⚠).
"""
import re, time
from docx import Document
from docx.shared import Pt, Cm, RGBColor

COMPANY = "ООО «Капитан Сэм»"
from answer import FORMAL   # формальные доводы («не распространяется», «не считается дефектом») — помечаются в черновике
GREY, RED = RGBColor(0x40, 0x40, 0x40), RGBColor(0xC0, 0x00, 0x00)

def restore(t, mapping):
    return re.sub(r"<[А-ЯA-Z]+_\d+>", lambda m: (mapping or {}).get(m.group(0), m.group(0)), t or "")

def calc_marker(t): return (t or "").replace("[Р]", "[расчёт]")

def _doc(size=10.5):
    d = Document()
    for s in d.sections: s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(size)
    return d

def _labelled(d, label, text, bold_label=True, italic=False, color=None):
    p = d.add_paragraph(); r = p.add_run(label); r.bold = bold_label
    r2 = p.add_run(text); r2.italic = italic
    if color: r2.font.color.rgb = color
    return p

def _structure(d, head, reasoning, basis, calc, action, clarify, show_calc, head_label="Вывод: "):
    """head — первая строка: покупателю «Коротко: <фраза по сути>», инженеру «Вывод: <тип вердикта>» (12.10.2026)."""
    if head: _labelled(d, head_label, calc_marker(head)).runs[1].bold = True
    if reasoning: d.add_paragraph(calc_marker(reasoning))
    if basis:
        d.add_paragraph().add_run("Основание:").bold = True
        for k, b in enumerate(basis, 1):
            p = d.add_paragraph(style="List Bullet"); p.add_run(f"[{k}] {b['label']}: ")
            q = p.add_run(f"«{b['quote']}»"); q.font.color.rgb = GREY
    if calc and show_calc:
        d.add_paragraph().add_run("Расчёт:").bold = True
        for l in calc: d.add_paragraph(l, style="List Bullet")
    if action: _labelled(d, "Что мы делаем: ", calc_marker(action))
    if clarify: _labelled(d, "Что нужно уточнить: ", calc_marker(clarify))

def render(outs, path, mapping=None, numbers=None, mark=None):
    """Ответ покупателю. numbers игнорируются: нумерация сквозная (исходный номер — только в черновике инженеру).
    mark — пометка под заголовком (12.10.2026: «по проекту (КП от …)» при повторном разборе с КП)."""
    d = _doc()
    p = d.add_paragraph(); r = p.add_run(COMPANY); r.bold = True; r.font.size = Pt(13)
    d.add_paragraph(time.strftime("%d.%m.%Y"))
    h = d.add_paragraph(); r = h.add_run("Ответ на обращение по проекту дома"); r.bold = True; r.font.size = Pt(14)
    if mark: d.add_paragraph().add_run(mark).bold = True
    d.add_paragraph("Благодарим Вас за внимательное отношение к проекту. Ниже — наш ответ по каждому пункту Вашего обращения.")
    for i, o in enumerate(outs, 1):
        m = o.get("mapping") or mapping
        d.add_paragraph().add_run(f"{i}.").bold = True
        _labelled(d, "Ваше замечание: ", restore(o["claim"], m), bold_label=False, italic=True).runs[0].italic = True
        used_calc = "[Р]" in " ".join(str(o.get(k) or "") for k in ("reasoning", "action", "clarify"))
        _structure(d, o.get("short"), o.get("reasoning"), o.get("basis"), o.get("calc"), o.get("action"), o.get("clarify"),
                   used_calc, head_label="Коротко: ")
    d.add_paragraph()
    d.add_paragraph("Если останутся вопросы — мы на связи и готовы обсудить их со специалистом.")
    d.add_paragraph(f"С уважением,\n{COMPANY}")
    d.save(path); return path

TYPES = {1: "1 — подтверждается нормативом", 2: "2 — опровергается нормативом / стандартом компании / практикой",
         3: "3 — встречный вопрос", 4: "4 — вне зоны ответственности", 5: "5 — основания не найдено"}

def render_engineer(items, path, mapping=None):
    """Черновик для инженера (ВНУТРЕННИЙ). items — [(номер пункта в обращении, out)]."""
    d = _doc(10)
    p = d.add_paragraph(); r = p.add_run("ВНУТРЕННИЙ ДОКУМЕНТ — покупателю не отправлять"); r.bold = True; r.font.color.rgb = RED
    h = d.add_paragraph(); r = h.add_run("Черновик для инженера"); r.bold = True; r.font.size = Pt(14)
    d.add_paragraph(f"{COMPANY} · {time.strftime('%d.%m.%Y')}. Пункты обращения, которые система не выдала покупателю. "
                    "По каждому — заготовка ответа; ⚠ — что проверка не подтвердила. Проверьте, поправьте и отправьте.")
    for n, o in items:
        i = o["internal"]; m = o.get("mapping") or mapping
        d.add_paragraph().add_run(f"Пункт {n} обращения").bold = True
        _labelled(d, "Претензия: ", restore(o["claim"], m))
        v, rv = i.get("verdict"), i.get("raw_verdict")
        vtxt = TYPES[v] if v in TYPES else (f"{TYPES[rv]} (ответ заблокирован проверкой)" if rv in TYPES else "не определён")
        _labelled(d, "Вероятный вердикт системы: ", vtxt)
        _labelled(d, "Почему не выдано покупателю: ", i.get("why_human") or "—", color=RED)
        if i.get("need_project"): _labelled(d, "Нужны данные проекта: ", i["need_project"], color=RED)
        if i.get("fear_hint"): _labelled(d, "Подсказка: ", i["fear_hint"], color=RED)
        d.add_paragraph().add_run("Что нашла система:").bold = True
        for f in i.get("found") or []:
            p = d.add_paragraph(style="List Bullet"); p.add_run(f["label"]).bold = True
            p.add_run(f" — {f['kind']}" + (f"; ⚠ {f['mark']}" if f.get("mark") else "") + (": " if f.get("quote") else ""))
            if f.get("quote"): q = p.add_run(f"«{f['quote']}»"); q.font.color.rgb = GREY
        dr = i.get("draft") or {}
        if any(dr.get(k) for k in ("conclusion", "reasoning", "action", "clarify")):
            d.add_paragraph().add_run("Заготовка ответа (по структуре ответа покупателю):").bold = True
            basis = [{"label": b.get("label_engineer", b["label"]), "quote": b["quote"] + (f"  ⚠ {b['mark']}" if b.get("mark") else "")}
                     for b in i.get("draft_basis") or []]
            reasoning = dr.get("reasoning")
            if i.get("fear") and reasoning:   # формальный довод оставляем как найденный, но с пометкой
                reasoning = " ".join(x + (" [⚠ формальный довод — для покупателя не годится как главный]" if FORMAL.search(x) else "")
                                     for x in re.split(r"(?<=[.!?])\s+(?=[А-ЯЁA-Z«\d])", reasoning))
            _structure(d, dr.get("conclusion"), reasoning, basis, o.get("calc"), dr.get("action"), dr.get("clarify"), bool(o.get("calc")))
        probs = i.get("problems") or []
        if probs:
            d.add_paragraph().add_run("⚠ Не подтверждено проверкой:").bold = True
            for pr in probs:
                p = d.add_paragraph(style="List Bullet"); r = p.add_run(f"{pr['what']}"); r.font.color.rgb = RED
                if pr.get("sentence") and pr["sentence"] not in ("вывод", "текст"): p.add_run(f" — «{calc_marker(pr['sentence'])}»")
    d.save(path); return path
