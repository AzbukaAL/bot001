"""Черновик для инженера: build/kamenskoe_draft_engineer.docx — 22 претензии Каменского (вкл. K041), ответ системы (прогон №8)
рядом с ответом продажника, пустые колонки «Верно? / Правка инженера». Без персональных данных (обезличено кодом).
06.10.2026: эталон v1 — позиция продаж, не эталон; эталон размечает инженер. P5–P10 в черновик не входят.
Система v1 текст ответа клиенту не формирует — в черновике её обоснование и встречный вопрос, как есть. Новых вызовов модели нет."""
import json, sys
sys.path.insert(0, "/opt/captain-sam/src")
from anonymize import anonymize
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = "/opt/captain-sam/build"
ET = {e["id"]: e for e in json.load(open(f"{B}/etalon.json", encoding="utf-8")) if e["id"].startswith("K")}
R = {r["id"]: r for r in json.load(open(f"{B}/run08.json", encoding="utf-8")) if r["id"].startswith("K")}
IDX = json.load(open(f"{B}/normy_index.json", encoding="utf-8"))
assert sorted(ET) == sorted(R) and len(R) == 22, (len(ET), len(R))
TYPES = {1: "1 — подтверждается нормативом", 2: "2 — опровергается нормативом / регламентом / практикой", 3: "3 — встречный вопрос",
         4: "4 — вне зоны ответственности", 5: "5 — основания не найдено"}
HOLD = "[цитата не выдаётся до сверки с оригиналом"
# 06.10.2026, решение Вадима: строки, где ошибка системы вероятна, — цветом и пометкой (сомнение одной фразой)
DOUBT = {
    "K020": "Проверьте особо: обоснование описывает крепление к винтовым сваям; какие сваи на объекте — ЖБ или винтовые?",
    "K022": "Проверьте особо: п. 6.17.7 СП 14 относится к фундаментам под сейсмоизоляторами — применим ли он к свайному фундаменту без сейсмоизоляции?",
    "K030": "Проверьте особо: возможна ошибка системы — фраза регламента общая («соблюдаются требования производителя») — на способ крепления в верх волны она не отвечает.",
    "K041": "Проверьте особо: возможна ошибка системы — границы ответственности в процитированных фрагментах нет — вывод «выбор материала — решение заказчика» сделала модель.",
}
# Претензия без обрывков исходной таблицы — как в build/etalon_soglasovanie.docx (src/etalon_docx.py, CLAIM_BRIEF).
# K032 там обрезан до «…», поэтому здесь полный текст без висящего «– )». K034 — без реплики покупателя.
CLAIM_CLEAN = {
    "K023": "Анкеровка каркаса к фундаменту — удержание от опрокидывания.",
    "K032": "Для ветрового района III с открытым рельефом нужна кровля с надёжным креплением — металлочерепица (саморезы в каждую волну), "
            "фальцевая кровля или мягкая (битумная) черепица на сплошную обрешётку.",
    "K057": "Черновой пол OSB 9 мм — тонко, прогибается при шаге лаг 600 мм.",
    "K059": "Брус 38×1900 мм в межэтажном перекрытии — опечатка, нужно уточнить.",
    "K060": "Чистовой пол OSB 22 мм — нужен финишный слой.",
    "K061": "Входная дверь (КНР) без терморазрыва — зимой при −7 °C и ветре будет промерзать.",
    "K062": "Бойлер 50 литров — мало для двухэтажного дома.",
}
CLAIM_CUT = {"K034": (" (я впервые стакиваюсь с таким строите—", " —"),   # реплика покупателя внутри текста претензии
             "K041": ("4. ", "")}                                        # номер пункта из исходного разбора

def claim(k, t):
    if k in CLAIM_CLEAN: return CLAIM_CLEAN[k]
    if k in CLAIM_CUT:
        cut, new = CLAIM_CUT[k]; assert cut in t, k; t = t.replace(cut, new, 1)
    return t

def A(t): return anonymize(" ".join(str(t).split()))[0]

def ocr_text(ref):
    """Текст пункта из индекса норм для удержанной цитаты (OCR / под изменениями) — инженеру, с пометкой «сверить»."""
    d, _, p = ref.partition("|п.")
    v = IDX.get(d, {}).get("points", {}).get(p)
    return v["text"][:300] if v else None

def quotes(r):
    out = []
    if r.get("evidence"): out.append(f"Доказательство (дословно из источника): «{r['evidence']}»")
    for ref, q in r["quotes"]:
        if q.startswith(HOLD):
            t = ocr_text(ref); why = q[len(HOLD):].strip(" :]")
            out.append(f"⚠ СВЕРИТЬ ({why or 'OCR'}): «{t}…» — {ref}" if t else f"⚠ цитата удержана до сверки, текста нет — {ref}")
        else:
            out.append(f"«{q}…» — {ref}")
    return "\n".join(out) or "—"

def system(r):
    if r["status"] != "ok":
        return "ЗАБЛОКИРОВАНО", f"Система ответ не выдала: {r['status'].replace('ОТВЕТ НЕ ВЫДАЁТСЯ: ', '')}."
    txt = f"Обоснование системы: {r['reason']}"
    if r.get("counter_question"): txt += f"\nВстречный вопрос клиенту: {r['counter_question']}"
    return TYPES.get(r["system"], f"вне пяти типов ({r['system']})"), txt

def sales(e):
    v = TYPES.get(e.get("verdict")) or "вердикт не дан"
    return f"[{v}]\n{e['answer']}"

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_); tcPr.append(s)

d = Document()
sec = d.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, m, Cm(1.2))
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(8)

d.add_heading("Каменское: ответы системы на проверку инженеру (черновик)", level=1)
for t in [
    "Что это. 22 претензии из разбора по с. Каменское. Для каждой — ответ системы (прогон №8 от 01.10.2026) и, для сравнения, "
    "ответ продажника. Ответы продажников эталоном не считаются. Эталон — Ваша разметка в двух последних колонках.",
    "Что заполнить. «Верно?» — да / нет / частично. «Правка инженера» — верный вердикт (один из пяти типов) и, если нужно, "
    "верное основание: норма с пунктом или положение регламента.",
    "Пять типов вердикта: 1 — подтверждается нормативом; 2 — опровергается нормативом, регламентом или практикой; "
    "3 — встречный вопрос; 4 — вне зоны ответственности; 5 — основания не найдено.",
    "Важно. Текста ответа клиенту система этой версии не пишет — показаны её обоснование и встречный вопрос как есть. "
    "«ЗАБЛОКИРОВАНО» — система ответ не выдала (проверка не пропустила); ответ за неё не дописан. "
    "«СВЕРИТЬ» — текст пункта распознан со скана или пункт затронут изменениями: сверить с оригиналом. "
    "СК-… — положения регламента компании (стандарт компании), не норматив: СК-ЧБ — файл «Ответы для чат-бота», "
    "СК-ВО — файл «Вопросы и ответы», § — номер раздела в файле. Строки, выделенные оранжевым, — «Проверьте особо». "
    "Персональные данные заменены метками.",
]: d.add_paragraph(A(t))

HDR = ["№", "Претензия клиента", "Вердикт системы", "Ответ системы", "Ссылки", "Цитаты", "Ответ продажника (для сравнения)",
       "Верно?", "Правка инженера"]
W = [1.1, 3.6, 2.2, 4.6, 2.0, 4.6, 4.0, 1.5, 3.4]
t = d.add_table(rows=1, cols=len(HDR)); t.style = "Table Grid"
for i, h in enumerate(HDR):
    c = t.rows[0].cells[i]; c.text = h; c.paragraphs[0].runs[0].bold = True; shade(c, "D9D9D9")
blocked = 0
for k in sorted(R):
    r, e = R[k], ET[k]; v, txt = system(r); blocked += v == "ЗАБЛОКИРОВАНО"
    ok = r["status"] == "ok"   # у заблокированного ответа ссылки модели проверку не прошли — не показываем
    row = [k, claim(k, e["claim"]), v, txt, ("; ".join(r.get("cite") or []) or "—") if ok else "—", quotes(r) if ok else "—", sales(e), "", ""]
    cells = t.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = A(val) if i not in (3, 5, 6) else "\n".join(A(x) for x in str(val).split("\n"))
    if k in DOUBT:
        cells[3].text = A(DOUBT[k]) + "\n" + cells[3].text
        cells[3].paragraphs[0].runs[0].bold = True
        for i in range(7): shade(cells[i], "FCE4D6")
    if v == "ЗАБЛОКИРОВАНО": shade(cells[2], "F4CCCC")
    shade(cells[7], "FFF2CC"); shade(cells[8], "FFF2CC")
for row in t.rows:
    for i, w in enumerate(W): row.cells[i].width = Cm(w)

out = f"{B}/kamenskoe_draft_engineer.docx"; d.save(out)

# контроль: обезличиватель по готовому тексту ничего нового не находит; обрывки исходной таблицы и реплика покупателя не остались
JUNK = ["| Акту", "Явная", "– .", "– )", "стакиваюсь"]
D = Document(out)
full = "\n".join(p.text for p in D.paragraphs) + "\n" + "\n".join(c.text for row in D.tables[0].rows for c in row.cells)
left = anonymize(full)[1]
print(f"строк: {len(D.tables[0].rows) - 1}, заблокировано: {blocked}, «проверьте особо»: {full.count('Проверьте особо: ')}, удержанных цитат: {full.count('СВЕРИТЬ (')}, "
      f"ПДн после обезличивания: {len(left)}")
for j in JUNK:
    hits = [row.cells[0].text for row in D.tables[0].rows[1:] if any(j in c.text for c in row.cells)] + (["шапка"] if j in "\n".join(p.text for p in D.paragraphs) else [])
    print(f"  «{j}»: {hits or 'нет'}")
