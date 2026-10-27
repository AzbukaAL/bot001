"""Претензия файлом (20.10.2026, решение Вадима): PDF / DOCX / ODT → текст → пункты по разделам.
Разбор нейросети покупателя приходит документом («1. Что предлагает подрядчик … 3.1. Фундамент …»): пункт — раздел
без подразделов (заголовок + текст); вступление раздела с подразделами — отдельным пунктом, только если оно по существу;
текст до первого раздела («Предварительное заключение») — тоже пункт. Обезличивание — дальше, в answer.analyze."""
import html, re, subprocess, zipfile

def extract_text(path):
    p = path.lower()
    if p.endswith(".pdf"):   # без -layout: порядок чтения, таблицы не разъезжаются по колонкам
        return subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True, check=True).stdout
    if p.endswith(".docx"):
        from docx import Document
        d = Document(path)
        return "\n".join([x.text for x in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    if p.endswith(".odt"):   # как pdn_check: zip + content.xml
        x = zipfile.ZipFile(path).read("content.xml").decode("utf-8")
        x = re.sub(r"<text:(?:p|h)\b[^>]*>|<text:line-break/>|<table:table-cell\b[^>]*>", "\n", x)
        x = re.sub(r"<text:s(?: text:c=\"(\d+)\")?/>", lambda m: " " * int(m.group(1) or 1), x)
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<text:tab/>", "\t", x)))
    raise ValueError(f"формат не поддерживается: {path}")

HEAD = re.compile(r"^\s*(\d{1,2}(?:\.\d{1,2}){0,2})\.?\s+([А-ЯЁA-Z«\"].{2,150})$")
INTRO_MIN, PRE_MIN, MAX_LEN = 200, 80, 2500

def split_sections(text):
    """→ [(заголовок, текст)] в порядке документа."""
    secs, pre, cur = [], [], None
    for line in text.splitlines():
        m = HEAD.match(line)
        if m:
            cur = {"num": m.group(1), "title": m.group(2).strip(), "body": []}; secs.append(cur); continue
        if line.strip(): (cur["body"] if cur else pre).append(line.strip())
    nums = [s["num"] for s in secs]
    out = []
    pre_t = " ".join(pre)
    if len(pre_t) >= PRE_MIN: out.append(("Общее заключение", pre_t))
    for s in secs:
        body = " ".join(s["body"]); parent = any(n.startswith(s["num"] + ".") for n in nums)
        if not parent or len(body) >= INTRO_MIN: out.append((f"{s['num']}. {s['title']}", body))
    return out

def points(text):
    """Пункты для разбора: «заголовок: текст», не длиннее MAX_LEN (обрезка — по границе предложения)."""
    res = []
    for title, body in split_sections(text):
        t = f"{title}: {body}" if body else title
        if len(t) > MAX_LEN:
            cut = t[:MAX_LEN]; k = max(cut.rfind(". "), cut.rfind("! "), cut.rfind("? "))
            t = cut[:k + 1] if k > MAX_LEN // 2 else cut
        res.append(" ".join(t.split()))
    return res
