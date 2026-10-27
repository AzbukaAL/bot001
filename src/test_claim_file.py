"""Претензия файлом, кнопки, подсказка, очередь уведомлений (20.10.2026, решение Вадима) — бот MAX на поддельном API,
без вызовов модели (разбор подменён).  python3 src/test_claim_file.py"""
import datetime, json, os, re, sys, tempfile, threading, urllib.error, zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ALEXEY = "/opt/captain-sam/build/mvp/bot/kp/kp_max_20261020_151943.pdf"   # реальная претензия 20.10 (ChatGPT, 52Р5)
KP = f"{ROOT}/materials/objects/test1-kp.pdf"
FILES = {"claim.pdf": ALEXEY, "test1-kp.pdf": KP}

LOG, QUEUE = [], []
class H(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def _send(self, code, obj=None, raw=None):
        b = raw if raw is not None else json.dumps(obj).encode(); self.send_response(code)
        self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)
    def do_GET(self):
        LOG.append(("GET", self.path, b""))
        if self.path.startswith("/files/"):
            n = self.path.split("/files/")[1]
            return self._send(200, raw=open(FILES[n], "rb").read() if n in FILES else open(EXTRA[n], "rb").read())
        return self._send(200, {"updates": QUEUE.pop(0) if QUEUE else [], "marker": 42})
    def do_POST(self):
        body = self.rfile.read(int(self.headers.get("Content-Length", 0))); LOG.append(("POST", self.path, body))
        if self.path.startswith("/bot"): return self._send(200, {"ok": True})
        if self.path.startswith("/uploads"): return self._send(200, {"url": f"http://127.0.0.1:{PORT}/upload-here"})
        if self.path.startswith("/upload-here"): return self._send(200, {"token": "T"})
        self._send(200, {"message": {}, "success": True})
EXTRA = {}
srv = HTTPServer(("127.0.0.1", 0), H); PORT = srv.server_port
threading.Thread(target=srv.serve_forever, daemon=True).start()
tmp = tempfile.mkdtemp(); open(f"{tmp}/notify", "w").write("N")
os.environ.update(CS_MAX_API=f"http://127.0.0.1:{PORT}", CS_TG_API=f"http://127.0.0.1:{PORT}", CS_TG_NOTIFY=f"{tmp}/notify",
                  CS_MAX_ACCESS=f"{tmp}/access.json", CS_MAX_KP=f"{tmp}/kp", CS_NOTIFY_FAILED=f"{tmp}/failed.jsonl")
import bot_max as B, claim_doc as C, daily_cost as DC
B.ALLOWED = {111}; B.STATE = f"{tmp}/marker"; B.OUT = tmp; B.time.sleep = lambda s: None
CALLS = []
def fake_handle(text, docs=None, redo=False, claims=None, first=1, part=None):
    CALLS.append({"text": text, "docs": list(docs or []), "claims": list(claims or []), "first": first, "part": part})
    p = f"{tmp}/x_{len(CALLS)}.docx"; open(p, "wb").write(b"PK"); n = len(claims) if claims else 1
    return [p], (f"Часть {part[0]} из {part[1]}. " if part else "") + f"Пунктов {n}.", [{"route": "engineer"}] * n
B.handle = fake_handle

ok = 0; fail = []
def check(name, cond):
    global ok
    if cond: ok += 1; print("✅", name)
    else: fail.append(name); print("❌", name)
def mc(uid, text="", att=None):
    body = {"text": text}
    if att: body["attachments"] = [att]
    return {"update_type": "message_created", "message": {"sender": {"user_id": uid, "name": "Олег"}, "body": body}}
def fileatt(name): return {"type": "file", "filename": name, "payload": {"url": f"http://127.0.0.1:{PORT}/files/{name}"}}
def cbk(uid, payload, cid="cb1"): return {"update_type": "message_callback", "callback": {"callback_id": cid, "payload": payload, "user": {"user_id": uid}}}
marker = None
def step(*ups):
    global marker; QUEUE.append(list(ups)); marker = B.poll_once("TOK", marker)
def posts(prefix="/messages"): return [json.loads(x[2]) for x in LOG if x[0] == "POST" and x[1].startswith(prefix) and x[2]]
def texts(): return [p.get("text") or "" for p in posts()]

# 1. PDF с претензией 52Р5 → не КП → кнопки
if not os.path.exists(ALEXEY): print("⚠ нет файла претензии 20.10 — сценарий 52Р5 пропущен")
else:
    n0 = len(LOG); step(mc(111, "", fileatt("claim.pdf")))
    kb = [p for p in posts() if any(a.get("type") == "inline_keyboard" for a in p.get("attachments") or [])]
    btns = kb[-1]["attachments"][0]["payload"]["buttons"][0] if kb else []
    check("52Р5: спросил кнопками", [b["text"] for b in btns] == ["Это претензия", "Это данные проекта"] and all(b["type"] == "callback" for b in btns))
    check("52Р5: до кнопки разбора нет и в данные проекта не попал", not CALLS and 111 not in B.PENDING and 111 not in B.REDO)
    key = btns[0]["payload"].split(":")[1]
    step(cbk(111, f"claim:{key}"))
    ans = [x for x in LOG if x[0] == "POST" and x[1].startswith("/answers")]
    check("кнопка: ответ на callback с callback_id, сообщение заменено", ans and "callback_id=cb1" in ans[-1][1] and json.loads(ans[-1][2])["message"]["text"].endswith("Разбираю."))
    pts = C.points(C.extract_text(ALEXEY))
    check("52Р5: разделов 12 (общее заключение, 1, 2, 3.1–3.4, 4.1–4.3, 5, 6)", len(pts) == 12 and pts[3].startswith("3.1.") and pts[-1].startswith("6."))
    check("52Р5: разобраны ВСЕ пункты частями 10 + 2", [len(c["claims"]) for c in CALLS] == [10, 2] and sum((c["claims"] for c in CALLS), []) == pts)
    check("52Р5: нумерация сквозная, части помечены", [c["first"] for c in CALLS] == [1, 11] and [c["part"] for c in CALLS] == [(1, 2), (2, 2)])
    check("52Р5: сотруднику «пунктов 12, частями (2)»", any("как претензию: пунктов 12, разберу частями (2)" in t for t in texts()))
    check("52Р5: повтор «в течение 30 минут» к претензии-файлу не привязан", 111 not in B.ANSWERED)
    check("52Р5: уведомление Вадиму «Претензия файлом»", any("Претензия файлом" in (p.get("text") or "") and "Пунктов: 12" in p["text"] for p in posts("/bot")))
    step(cbk(111, f"claim:{key}", "cb2"))
    check("повторное нажатие — «кнопка устарела», без разбора", len(CALLS) == 2 and "устарела" in json.loads([x for x in LOG if x[1].startswith("/answers")][-1][2])["message"]["text"])
    # кнопка «данные проекта» → прежний путь
    CALLS.clear(); step(mc(111, "", fileatt("claim.pdf")))
    key2 = [p for p in posts() if p.get("attachments")][-1]["attachments"][0]["payload"]["buttons"][0][1]["payload"].split(":")[1]
    step(cbk(111, f"data:{key2}", "cb3"))
    check("кнопка «данные проекта» → файл ждёт текст (прежний путь)", not CALLS and 111 in B.PENDING and len(B.PENDING[111]["docs"]) == 1)
    B.PENDING.clear()
    # текст вместе с нераспознанным файлом → разбор текста + подсказка про кнопки
    step(mc(111, "Кровля из ондулина протечёт на стыках", fileatt("claim.pdf")))
    check("текст + нераспознанный файл: текст разобран без файла", CALLS and CALLS[-1]["text"].startswith("Кровля") and not CALLS[-1]["docs"])
    check("в ответе — «если это претензия — нажмите кнопку или пришлите текстом»", any(B.CLAIM_BTN_HINT in t and "ждёт ответа на кнопки" in t for t in texts()))
    # кнопку не нажали → отпущен с уведомлением
    B.KP_HOLD = 0; step(); B.KP_HOLD = 1800
    check("кнопку не нажали — файл отпущен, Вадиму уведомление", not B.ASKS and any("кнопку не нажали" in (p.get("text") or "") for p in posts("/bot")))

# 2. «1», «ок», «привет» → подсказка; короткое со строительным словом — разбор
CALLS.clear(); n0 = len(texts())
for t in ("1", "ок", "привет", "Спасибо!"): step(mc(111, t))
check("«1», «ок», «привет», «Спасибо!» → подсказка, без разбора и черновика", not CALLS and sum(B.SHORT_HINT == t for t in texts()[n0:]) == 4)
step(mc(111, "Течёт кровля"))
check("«Течёт кровля» (короткое, но по делу) → разбор", len(CALLS) == 1)

# 3. КП «Тест 1» → распознано → как раньше, без кнопок
CALLS.clear(); B.ANSWERED.clear(); B.PENDING.clear(); kb0 = sum(1 for p in posts() if p.get("attachments"))
step(mc(111, "", fileatt("test1-kp.pdf")))
check("КП «Тест 1» → данные проекта, без кнопок", 111 in B.PENDING and sum(1 for p in posts() if p.get("attachments")) == kb0
      and any(t.startswith("Получил: КП от") for t in texts()))
B.PENDING.clear()

# 4. DOCX и ODT — претензия файлом
from docx import Document
d = Document(); d.add_paragraph("Разбор проекта нейросетью. Общее замечание по всем узлам дома, требует ответа подрядчика по существу вопроса.")
d.add_paragraph("1. Кровля"); d.add_paragraph("Ондулин протечёт на стыках листов."); d.add_paragraph("2. Стены"); d.add_paragraph("OSB 9 мм накопит влагу.")
d.save(f"{tmp}/c.docx"); EXTRA["c.docx"] = f"{tmp}/c.docx"
check("DOCX: пункты по разделам", C.points(C.extract_text(f"{tmp}/c.docx")) == [
    "Общее заключение: Разбор проекта нейросетью. Общее замечание по всем узлам дома, требует ответа подрядчика по существу вопроса.",
    "1. Кровля: Ондулин протечёт на стыках листов.", "2. Стены: OSB 9 мм накопит влагу."])
odt = f"{tmp}/c.odt"
with zipfile.ZipFile(odt, "w") as z:
    z.writestr("content.xml", '<office:document-content xmlns:office="o" xmlns:text="t"><office:body><office:text>'
               '<text:p>1. Окна</text:p><text:p>Стеклопакет 4-16-4 будет потеть.</text:p><text:p>2. Дверь</text:p><text:p>Промёрзнет.</text:p>'
               '</office:text></office:body></office:document-content>')
check("ODT: пункты по разделам", C.points(C.extract_text(odt)) == ["1. Окна: Стеклопакет 4-16-4 будет потеть.", "2. Дверь: Промёрзнет."])
CALLS.clear(); step(mc(111, "", fileatt("c.docx")))
k3 = [p for p in posts() if p.get("attachments")][-1]["attachments"][0]["payload"]["buttons"][0][0]["payload"]
step(cbk(111, k3, "cb4"))
check("DOCX через бота: кнопка → разбор 3 пунктов", len(CALLS) == 1 and len(CALLS[0]["claims"]) == 3 and CALLS[0]["part"] is None)

# 4а. Скан без распознавания в кэше — сразу «Распознаю скан — это 1–2 минуты» (21.10)
CALLS.clear(); n0 = len(texts()); B.scan_pending = lambda paths: True
B.process("TOK", 111, "Кровля из ондулина протечёт на стыках", [f"{tmp}/c.docx"])
t_new = texts()[n0:]
check("скан: сначала «Принято…», затем «Распознаю скан — это 1–2 минуты», потом разбор", len(t_new) >= 2 and t_new[1] == B.OCR_NOTE and CALLS)
CALLS.clear(); n0 = len(texts()); B.scan_pending = lambda paths: False
B.process("TOK", 111, "Кровля из ондулина протечёт на стыках", [f"{tmp}/c.docx"])
check("не скан / уже распознан — без сообщения о распознавании", B.OCR_NOTE not in texts()[n0:])

# 5. Уведомления: осечка → повтор; 3 осечки → журнал неотправленных и вечерняя строка
real = B.tg_send; dead = {"n": 0}
def tg_dead(body): dead["n"] += 1; raise urllib.error.URLError("нет связи")
B.tg_send = tg_dead; B.notify("тест: уведомление 1")
check("осечка → в очереди, попытка 1", len(B.NOTIFY_Q) == 1 and B.NOTIFY_Q[0]["tries"] == 1)
B.NOTIFY_PAUSE = 0; B.NOTIFY_Q[0]["next"] = 0
B.tg_send = real; B.notify_retry()
check("повтор удался → очередь пуста", not B.NOTIFY_Q and any("(с задержкой)" in (p.get("text") or "") for p in posts("/bot")))
B.tg_send = tg_dead; B.notify("тест: уведомление 2")
for _ in range(3): B.notify_retry()
rows = [json.loads(l) for l in open(f"{tmp}/failed.jsonl", encoding="utf-8")]
check("3 осечки → журнал неотправленных, очередь пуста", not B.NOTIFY_Q and len(rows) == 1 and rows[0]["head"] == "тест: уведомление 2" and dead["n"] == 4)
B.tg_send = real
now = datetime.datetime.strptime(rows[0]["ts"], "%Y-%m-%d %H:%M:%S") + datetime.timedelta(minutes=1)
open(f"{tmp}/calls.jsonl", "w").write("")
check("вечерняя строка: «уведомлений не ушло: 1»", "уведомлений не ушло: 1" in DC.line(now, f"{tmp}/calls.jsonl", tmp, f"{tmp}/failed.jsonl"))
check("вечерняя строка без журнала неотправленных — без лишнего", "уведомлений" not in DC.line(now, f"{tmp}/calls.jsonl", tmp, f"{tmp}/нет.jsonl"))
print(f"претензия файлом и кнопки: {ok}/{ok + len(fail)}", "" if not fail else f"ПРОВАЛ: {fail}")
sys.exit(1 if fail else 0)
